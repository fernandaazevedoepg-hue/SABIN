from datetime import datetime
from pathlib import Path
import os
import subprocess
import sys

import pandas as pd
from dotenv import load_dotenv
from sqlalchemy import create_engine, text
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table as ExcelTable
from openpyxl.worksheet.table import TableStyleInfo
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

load_dotenv()

BASE_DIR = Path(__file__).resolve().parent.parent
EXPORTS_DIR = BASE_DIR / "exports"
DATABASE_URL = os.getenv("DATABASE_URL")
if not DATABASE_URL:
    raise ValueError("DATABASE_URL não encontrada no ficheiro .env")

engine = create_engine(DATABASE_URL, pool_pre_ping=True)

AZUL_ESCURO = "1F3A5F"
AZUL_CLARO = "EAF1FF"
CINZA_CLARO = "F4F7FB"
CINZA_BORDA = "D8E0EA"
TEXTO = "172033"
BORDA = Side(style="thin", color=CINZA_BORDA)
FORMATO_EUROS = '#,##0.00 "€"' 


def ler_df(sql):
    with engine.connect() as connection:
        return pd.read_sql(text(sql), connection)


def carregar_kpis():
    linha = ler_df("""
        SELECT
            COALESCE(SUM(iv.subtotal), 0) AS receita_total,
            COUNT(DISTINCT v.id) AS numero_vendas,
            COALESCE(SUM(iv.quantidade), 0) AS unidades_vendidas
        FROM public.vendas v
        JOIN public.itens_venda iv ON iv.venda_id = v.id
        WHERE v.status = 'Concluída';
    """).iloc[0]

    receita = float(linha["receita_total"])
    vendas = int(linha["numero_vendas"])
    unidades = int(linha["unidades_vendidas"])
    return {
        "receita_total": receita,
        "numero_vendas": vendas,
        "unidades_vendidas": unidades,
        "ticket_medio": receita / vendas if vendas else 0,
    }


def carregar_metricas():
    try:
        dados = ler_df("SELECT * FROM dw.metricas_previsao LIMIT 1;")
        if dados.empty:
            return {}
        linha = dados.iloc[0]
        return {
            coluna: linha[coluna]
            for coluna in dados.columns
            if pd.notna(linha[coluna])
        }
    except Exception as erro:
        print("Erro ao carregar métricas:", erro)
        return {}


def carregar_vendas():
    return ler_df("""
        SELECT
            v.data_venda AS data,
            v.id AS venda_id,
            iv.id AS item_venda_id,
            COALESCE(c.nome_completo, 'Cliente não identificado') AS cliente,
            l.titulo AS livro,
            l.isbn,
            v.metodo_pagamento,
            iv.quantidade,
            iv.preco_unitario,
            iv.subtotal
        FROM public.vendas v
        JOIN public.itens_venda iv ON iv.venda_id = v.id
        JOIN public.livros l ON l.id = iv.livro_id
        LEFT JOIN public.clientes c ON c.id = v.cliente_id
        WHERE v.status = 'Concluída'
        ORDER BY v.data_venda, v.id, iv.id;
    """)


def carregar_livros():
    return ler_df("""
        SELECT
            l.id,
            l.titulo,
            l.isbn,
            l.editora,
            l.data_publicacao,
            l.preco_venda,
            l.estoque_atual AS stock_fisico,
            l.qtd_reservada AS reservado,
            GREATEST(l.estoque_atual - l.qtd_reservada, 0) AS stock_disponivel,
            COALESCE((
                SELECT STRING_AGG(a.nome, ', ' ORDER BY a.nome)
                FROM public.livro_autores la
                JOIN public.autores a ON a.id = la.autor_id
                WHERE la.livro_id = l.id
            ), '') AS autores,
            COALESCE((
                SELECT STRING_AGG(g.nome, ', ' ORDER BY g.nome)
                FROM public.livro_generos lg
                JOIN public.generos g ON g.id = lg.genero_id
                WHERE lg.livro_id = l.id
            ), '') AS generos
        FROM public.livros l
        ORDER BY l.titulo;
    """)


def carregar_clientes():
    return ler_df("""
        SELECT
            c.id,
            c.nome_completo,
            c.nif,
            c.telemovel,
            c.email,
            c.data_registo,
            COALESCE(compras.valor_total, 0) AS total_compras_valor,
            COALESCE(compras.numero_compras, 0) AS numero_compras,
            COALESCE(compras.unidades_compradas, 0) AS unidades_compradas
        FROM public.clientes c
        LEFT JOIN (
            SELECT
                v.cliente_id,
                COUNT(DISTINCT v.id) AS numero_compras,
                COALESCE(SUM(iv.quantidade), 0) AS unidades_compradas,
                COALESCE(SUM(iv.subtotal), 0) AS valor_total
            FROM public.vendas v
            JOIN public.itens_venda iv ON iv.venda_id = v.id
            WHERE v.status = 'Concluída'
              AND v.cliente_id IS NOT NULL
            GROUP BY v.cliente_id
        ) compras ON compras.cliente_id = c.id
        ORDER BY c.nome_completo;
    """)


def carregar_stock():
    return ler_df("""
        SELECT
            titulo,
            estoque_atual AS stock_fisico,
            qtd_reservada AS reservado,
            stock_disponivel,
            unidades_ultimos_3_meses,
            media_mensal_vendas,
            meses_cobertura,
            quantidade_recomendada,
            nivel,
            recomendacao
        FROM dw.alertas_stock
        ORDER BY
            CASE
                WHEN nivel = 'Crítico' THEN 1
                WHEN nivel = 'Atenção' THEN 2
                WHEN nivel = 'Normal' THEN 3
                ELSE 4
            END,
            meses_cobertura NULLS LAST,
            titulo;
    """)


def carregar_previsoes():
    return ler_df("""
        SELECT mes, receita_real, receita_prevista, tipo
        FROM dw.previsao_receita
        ORDER BY mes;
    """)


def dados_exportacao():
    return {
        "kpis": carregar_kpis(),
        "metricas": carregar_metricas(),
        "Vendas": carregar_vendas(),
        "Livros": carregar_livros(),
        "Clientes": carregar_clientes(),
        "Stock": carregar_stock(),
        "Previsoes": carregar_previsoes(),
    }


def valor_excel(valor):
    if valor is None:
        return None
    try:
        if pd.isna(valor):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(valor, pd.Timestamp):
        valor = valor.to_pydatetime()
    if isinstance(valor, datetime) and valor.tzinfo is not None:
        valor = valor.replace(tzinfo=None)
    return valor


def aplicar_borda(celula):
    celula.border = Border(left=BORDA, right=BORDA, top=BORDA, bottom=BORDA)


def titulo_folha(folha, ultima_coluna, titulo, descricao):
    folha.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ultima_coluna)
    celula_titulo = folha["A1"]
    celula_titulo.value = titulo
    celula_titulo.font = Font(bold=True, size=18, color="FFFFFF")
    celula_titulo.fill = PatternFill("solid", fgColor=AZUL_ESCURO)
    celula_titulo.alignment = Alignment(vertical="center")
    folha.row_dimensions[1].height = 30

    folha.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ultima_coluna)
    celula_descricao = folha["A2"]
    celula_descricao.value = (
        f"{descricao} | Gerado em {datetime.now().strftime('%d/%m/%Y às %H:%M')}"
    )
    celula_descricao.font = Font(italic=True, size=10, color="58677C")


def criar_resumo_excel(workbook, dados):
    folha = workbook.active
    folha.title = "Resumo"
    folha.sheet_view.showGridLines = False
    titulo_folha(folha, 4, "SABIN — Relatório de Gestão", "Resumo executivo da Bookmarked")

    kpis = dados["kpis"]
    metricas = dados["metricas"]
    stock = dados["Stock"]

    folha["A4"] = "Indicadores principais"
    folha["A4"].font = Font(bold=True, size=13, color=TEXTO)
    cards = [
        ("Receita Total", kpis["receita_total"], FORMATO_EUROS),
        ("Número de Vendas", kpis["numero_vendas"], "#,##0"),
        ("Unidades Vendidas", kpis["unidades_vendidas"], "#,##0"),
        ("Ticket Médio", kpis["ticket_medio"], FORMATO_EUROS),
    ]
    for coluna, (titulo, valor, formato) in enumerate(cards, 1):
        titulo_cell = folha.cell(5, coluna, titulo)
        valor_cell = folha.cell(6, coluna, valor)
        for celula in (titulo_cell, valor_cell):
            celula.fill = PatternFill("solid", fgColor=CINZA_CLARO)
            celula.alignment = Alignment(horizontal="center", vertical="center")
            aplicar_borda(celula)
        titulo_cell.font = Font(bold=True, size=9, color="5C6B80")
        valor_cell.font = Font(bold=True, size=14, color=AZUL_ESCURO)
        valor_cell.number_format = formato
    folha.row_dimensions[5].height = 24
    folha.row_dimensions[6].height = 32

    folha["A9"] = "Previsão e qualidade do modelo"
    folha["A9"].font = Font(bold=True, size=13, color=TEXTO)
    previsao = [
        ("Mês Previsto", metricas.get("proximo_mes"), "mm/yyyy"),
        ("Receita Prevista", metricas.get("receita_prevista_proximo_mes"), FORMATO_EUROS),
        ("Variação Prevista", metricas.get("variacao_prevista_percent"), '0.00"%"'),
        ("Tendência", metricas.get("tendencia"), "General"),
        ("Modelo", metricas.get("modelo"), "General"),
        ("MAE", metricas.get("mae"), FORMATO_EUROS),
        ("RMSE", metricas.get("rmse"), FORMATO_EUROS),
        ("R²", metricas.get("r2"), "0.0000"),
    ]
    for indice, (titulo, valor, formato) in enumerate(previsao):
        coluna = 1 if indice < 4 else 3
        linha = 10 + (indice if indice < 4 else indice - 4)
        titulo_cell = folha.cell(linha, coluna, titulo)
        valor_cell = folha.cell(linha, coluna + 1, valor_excel(valor))
        titulo_cell.font = Font(bold=True, color=TEXTO)
        titulo_cell.fill = PatternFill("solid", fgColor=AZUL_CLARO)
        valor_cell.fill = PatternFill("solid", fgColor="FFFFFF")
        valor_cell.number_format = formato
        aplicar_borda(titulo_cell)
        aplicar_borda(valor_cell)

    folha["A16"] = "Stock e reposição"
    folha["A16"].font = Font(bold=True, size=13, color=TEXTO)
    if stock.empty:
        criticos = atencao = repor = 0
    else:
        criticos = int((stock["nivel"] == "Crítico").sum())
        atencao = int((stock["nivel"] == "Atenção").sum())
        repor = int(pd.to_numeric(stock["quantidade_recomendada"], errors="coerce").fillna(0).sum())
    stock_cards = [
        ("Livros Analisados", len(stock)),
        ("Críticos", criticos),
        ("Em Atenção", atencao),
        ("Unidades a Repor", repor),
    ]
    for coluna, (titulo, valor) in enumerate(stock_cards, 1):
        titulo_cell = folha.cell(17, coluna, titulo)
        valor_cell = folha.cell(18, coluna, valor)
        for celula in (titulo_cell, valor_cell):
            celula.fill = PatternFill("solid", fgColor=CINZA_CLARO)
            celula.alignment = Alignment(horizontal="center")
            aplicar_borda(celula)
        titulo_cell.font = Font(bold=True, size=9, color="5C6B80")
        valor_cell.font = Font(bold=True, size=14, color=AZUL_ESCURO)
    for coluna in "ABCD":
        folha.column_dimensions[coluna].width = 24
    folha.freeze_panes = "A4"


def criar_folha_dados(workbook, nome, dataframe, mapa, titulo, descricao):
    folha = workbook.create_sheet(nome)
    folha.sheet_view.showGridLines = False
    dados = dataframe.rename(columns=mapa).copy()
    cabecalhos = list(dados.columns)
    ultima_coluna = max(1, len(cabecalhos))
    titulo_folha(folha, ultima_coluna, titulo, descricao)

    for coluna, cabecalho in enumerate(cabecalhos, 1):
        celula = folha.cell(4, coluna, cabecalho)
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor=AZUL_ESCURO)
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        aplicar_borda(celula)

    for linha_excel, linha in enumerate(dados.itertuples(index=False, name=None), 5):
        for coluna_excel, valor in enumerate(linha, 1):
            celula = folha.cell(linha_excel, coluna_excel, valor_excel(valor))
            celula.alignment = Alignment(vertical="top", wrap_text=True)
            aplicar_borda(celula)

    if not dados.empty and cabecalhos:
        referencia = f"A4:{get_column_letter(ultima_coluna)}{folha.max_row}"
        tabela = ExcelTable(displayName=f"Tabela{nome}", ref=referencia)
        tabela.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        folha.add_table(tabela)

    formatos = {
        "Data": "dd/mm/yyyy",
        "Data de Publicação": "dd/mm/yyyy",
        "Data de Registo": "dd/mm/yyyy",
        "Período": "mm/yyyy",
        "Preço Unitário (€)": FORMATO_EUROS,
        "Subtotal (€)": FORMATO_EUROS,
        "Preço de Venda (€)": FORMATO_EUROS,
        "Total Compras (€)": FORMATO_EUROS,
        "Receita Real (€)": FORMATO_EUROS,
        "Receita Prevista (€)": FORMATO_EUROS,
        "Média Mensal de Vendas": "0.00",
        "Cobertura (Meses)": "0.00",
    }
    for indice, cabecalho in enumerate(cabecalhos, 1):
        maior = len(str(cabecalho))
        for linha in range(5, folha.max_row + 1):
            celula = folha.cell(linha, indice)
            if cabecalho in formatos:
                celula.number_format = formatos[cabecalho]
            if celula.value is not None:
                maior = max(maior, len(str(celula.value)))
        folha.column_dimensions[get_column_letter(indice)].width = min(max(maior + 2, 12), 42)
    folha.freeze_panes = "A5"


def gerar_excel():
    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    dados = dados_exportacao()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    caminho = EXPORTS_DIR / f"SABIN_Relatorio_Gestao_{timestamp}.xlsx"
    workbook = Workbook()
    criar_resumo_excel(workbook, dados)

    configuracoes = {
        "Vendas": (
            {
                "data": "Data", "venda_id": "ID Venda", "item_venda_id": "ID Item",
                "cliente": "Cliente", "livro": "Livro", "isbn": "ISBN",
                "metodo_pagamento": "Método de Pagamento", "quantidade": "Quantidade",
                "preco_unitario": "Preço Unitário (€)", "subtotal": "Subtotal (€)",
            },
            "Vendas Concluídas",
            "Detalhe dos itens pertencentes às vendas concluídas.",
        ),
        "Livros": (
            {
                "id": "ID", "titulo": "Título", "isbn": "ISBN", "editora": "Editora",
                "data_publicacao": "Data de Publicação", "preco_venda": "Preço de Venda (€)",
                "stock_fisico": "Stock Físico", "reservado": "Reservado",
                "stock_disponivel": "Stock Disponível", "autores": "Autores", "generos": "Géneros",
            },
            "Catálogo de Livros",
            "Informação atual do catálogo e da disponibilidade de stock.",
        ),
        "Clientes": (
            {
                "id": "ID", "nome_completo": "Nome", "nif": "NIF",
                "telemovel": "Telemóvel", "email": "Email",
                "data_registo": "Data de Registo", "total_compras_valor": "Total Compras (€)",
                "numero_compras": "Número de Compras", "unidades_compradas": "Unidades Compradas",
            },
            "Clientes",
            "Compras distintas, unidades e montantes de vendas concluídas por cliente.",
        ),
        "Stock": (
            {
                "titulo": "Livro", "stock_fisico": "Stock Físico", "reservado": "Reservado",
                "stock_disponivel": "Stock Disponível",
                "unidades_ultimos_3_meses": "Unidades Últimos 3 Meses",
                "media_mensal_vendas": "Média Mensal de Vendas",
                "meses_cobertura": "Cobertura (Meses)",
                "quantidade_recomendada": "Quantidade a Repor",
                "nivel": "Nível", "recomendacao": "Recomendação",
            },
            "Stock e Reposição",
            "Indicadores utilizados para análise de cobertura e reposição.",
        ),
        "Previsoes": (
            {
                "mes": "Período", "receita_real": "Receita Real (€)",
                "receita_prevista": "Receita Prevista (€)", "tipo": "Tipo",
            },
            "Previsões de Receita",
            "Histórico observado e estimativas produzidas pelo modelo preditivo.",
        ),
    }

    for nome, (mapa, titulo, descricao) in configuracoes.items():
        criar_folha_dados(workbook, nome, dados[nome], mapa, titulo, descricao)
    workbook.save(caminho)
    return caminho


def euro(valor):
    if valor is None or pd.isna(valor):
        return "—"
    return f"{float(valor):,.2f} €".replace(",", "X").replace(".", ",").replace("X", " ")


def numero(valor):
    if valor is None or pd.isna(valor):
        return "—"
    return f"{int(valor):,}".replace(",", " ")


def decimal(valor, casas=2):
    if valor is None or pd.isna(valor):
        return "—"
    return f"{float(valor):.{casas}f}".replace(".", ",")


def mes(valor):
    if valor is None or pd.isna(valor):
        return "—"
    return pd.to_datetime(valor).strftime("%m/%Y")


def estilo_tabela_pdf():
    return TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D8E0EA")),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4.5),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
    ])


def rodape_pdf(canvas, documento):
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#D8E0EA"))
    canvas.line(documento.leftMargin, 12 * mm, A4[0] - documento.rightMargin, 12 * mm)
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#6B778C"))
    canvas.drawString(documento.leftMargin, 7.5 * mm, "SABIN — Business Intelligence")
    canvas.drawRightString(A4[0] - documento.rightMargin, 7.5 * mm, f"Página {documento.page}")
    canvas.restoreState()


def gerar_pdf():
    import math
    from html import escape

    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    dados = dados_exportacao()
    kpis = dados.get("kpis") or {}
    metricas = dados.get("metricas") or {}
    vendas = dados.get("Vendas")
    stock = dados.get("Stock")
    previsoes = dados.get("Previsoes")

    if not isinstance(vendas, pd.DataFrame):
        vendas = pd.DataFrame()
    if not isinstance(stock, pd.DataFrame):
        stock = pd.DataFrame()
    if not isinstance(previsoes, pd.DataFrame):
        previsoes = pd.DataFrame()

    def valor_numero(valor):
        try:
            resultado = float(valor)
            return resultado if math.isfinite(resultado) else None
        except (TypeError, ValueError):
            return None

    def inteiro(valor):
        numero_convertido = valor_numero(valor)
        return int(numero_convertido) if numero_convertido is not None else 0

    def seguro(valor):
        if valor is None:
            return "—"
        try:
            if pd.isna(valor):
                return "—"
        except (TypeError, ValueError):
            pass
        return escape(str(valor))

    def lista_titulos(tabela, limite=4):
        nomes = []
        for _, linha in tabela.head(limite).iterrows():
            titulo = str(linha.get("titulo", "Livro sem título"))
            qtd = inteiro(linha.get("quantidade_recomendada"))
            if qtd > 0:
                nomes.append(f"{titulo} ({qtd} un. sugeridas)")
            else:
                nomes.append(titulo)
        return "; ".join(nomes)

    receita = valor_numero(kpis.get("receita_total")) or 0.0
    n_vendas = inteiro(kpis.get("numero_vendas"))
    unidades = inteiro(kpis.get("unidades_vendidas"))
    ticket = valor_numero(kpis.get("ticket_medio")) or 0.0

    receita_prevista = valor_numero(metricas.get("receita_prevista_proximo_mes"))
    variacao = valor_numero(metricas.get("variacao_prevista_percent"))
    mae = valor_numero(metricas.get("mae"))
    rmse = valor_numero(metricas.get("rmse"))
    r2 = valor_numero(metricas.get("r2"))

    proximo_mes = metricas.get("proximo_mes")
    proximo_periodo = None
    try:
        if proximo_mes is not None and not pd.isna(proximo_mes):
            proximo_periodo = pd.to_datetime(proximo_mes).to_period("M")
    except (TypeError, ValueError):
        pass
    periodo = proximo_periodo.strftime("%m/%Y") if proximo_periodo is not None else "—"

    base_comparacao = None
    referencia = None
    if (
        proximo_periodo is not None
        and receita_prevista is not None
        and variacao is not None
        and {"mes", "receita_real"}.issubset(previsoes.columns)
    ):
        historico = previsoes.copy()
        historico["periodo"] = pd.to_datetime(historico["mes"], errors="coerce").dt.to_period("M")
        mes_anterior = proximo_periodo - 1
        correspondencias = historico[historico["periodo"].eq(mes_anterior)]
        for _, linha in correspondencias.iterrows():
            real_anterior = valor_numero(linha.get("receita_real"))
            if real_anterior is None or real_anterior <= 0:
                continue
            percentual = (receita_prevista / real_anterior - 1) * 100
            if abs(percentual - variacao) <= 0.5:
                base_comparacao = mes_anterior.strftime("%m/%Y")
                referencia = real_anterior
                break

    criticos = pd.DataFrame()
    atencao = pd.DataFrame()
    top_alertas = pd.DataFrame()
    total_repor = 0
    if not stock.empty and {"nivel", "quantidade_recomendada"}.issubset(stock.columns):
        criticos = stock[stock["nivel"].eq("Crítico")].copy()
        atencao = stock[stock["nivel"].eq("Atenção")].copy()
        total_repor = int(pd.to_numeric(
            stock["quantidade_recomendada"], errors="coerce"
        ).fillna(0).clip(lower=0).sum())
        top_alertas = stock[stock["nivel"].isin(["Crítico", "Atenção"])].head(12)

    top_livros = pd.DataFrame()
    if not vendas.empty and {"livro", "quantidade"}.issubset(vendas.columns):
        quantidades = vendas.copy()
        quantidades["quantidade"] = pd.to_numeric(
            quantidades["quantidade"], errors="coerce"
        ).fillna(0)
        top_livros = (
            quantidades.groupby("livro", as_index=False, dropna=False)["quantidade"]
            .sum().sort_values(["quantidade", "livro"], ascending=[False, True])
            .head(5)
        )
        top_livros = top_livros[top_livros["quantidade"] > 0]

    caminho = EXPORTS_DIR / (
        f"SABIN_Relatorio_Gestao_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    )
    documento = SimpleDocTemplate(
        str(caminho), pagesize=A4,
        leftMargin=16 * mm, rightMargin=16 * mm,
        topMargin=15 * mm, bottomMargin=18 * mm,
        title="SABIN - Relatório de Gestão", author="SABIN",
    )

    estilos = getSampleStyleSheet()
    estilo_titulo = ParagraphStyle(
        "TituloSABIN2026", parent=estilos["Title"], alignment=TA_CENTER,
        fontSize=21, leading=25, textColor=colors.HexColor("#172033"), spaceAfter=2 * mm,
    )
    estilo_subtitulo = ParagraphStyle(
        "SubtituloSABIN2026", parent=estilos["BodyText"], alignment=TA_CENTER,
        fontSize=9.5, leading=13, textColor=colors.HexColor("#5C6B80"),
    )
    estilo_secao = ParagraphStyle(
        "SecaoSABIN2026", parent=estilos["Heading2"], fontSize=12.5,
        leading=16, textColor=colors.HexColor("#172033"),
        spaceBefore=3.5 * mm, spaceAfter=2 * mm, keepWithNext=True,
    )
    estilo_normal = ParagraphStyle(
        "TextoSABIN2026", parent=estilos["BodyText"], fontSize=8.5,
        leading=12, textColor=colors.HexColor("#243044"), spaceAfter=2 * mm,
    )
    estilo_nota = ParagraphStyle(
        "NotaSABIN2026", parent=estilo_normal, fontSize=8.3,
        leading=11.5, textColor=colors.HexColor("#45617D"), spaceAfter=0,
    )
    estilo_label = ParagraphStyle(
        "LabelSABIN2026", parent=estilo_normal, alignment=TA_CENTER,
        fontName="Helvetica-Bold", fontSize=7.4, leading=10,
    )
    estilo_valor = ParagraphStyle(
        "ValorSABIN2026", parent=estilo_label,
        fontSize=12, leading=15, textColor=colors.HexColor("#1F3A5F"),
    )
    estilo_celula = ParagraphStyle(
        "CelulaSABIN2026", parent=estilo_normal, fontSize=7.9,
        leading=10, spaceAfter=0,
    )
    estilo_cabecalho = ParagraphStyle(
        "CabecalhoSABIN2026", parent=estilo_celula,
        textColor=colors.white, fontName="Helvetica-Bold", fontSize=7.6,
    )

    elementos = []

    def secao(titulo):
        elementos.append(Paragraph(seguro(titulo), estilo_secao))

    def paragrafo(texto):
        elementos.append(Paragraph(seguro(texto), estilo_normal))

    def bloco(texto, titulo="Interpretação"):
        tabela = Table(
            [[Paragraph(f"<b>{seguro(titulo)}:</b> {seguro(texto)}", estilo_nota)]],
            colWidths=[176 * mm], hAlign="LEFT", splitByRow=0,
        )
        tabela.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F4F7FB")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#D8E0EA")),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 7),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ]))
        elementos.extend([Spacer(1, 1.2 * mm), tabela, Spacer(1, 1.5 * mm)])

    def tabela_dados(cabecalhos, linhas, larguras):
        conteudo = [[Paragraph(seguro(c), estilo_cabecalho) for c in cabecalhos]]
        conteudo.extend(linhas)
        tabela = Table(conteudo, colWidths=larguras, repeatRows=1, hAlign="LEFT")
        tabela.setStyle(estilo_tabela_pdf())
        elementos.append(tabela)

    elementos.extend([
        Paragraph("SABIN", estilo_titulo),
        Paragraph("Relatório de Gestão da Bookmarked", estilo_subtitulo),
        Paragraph(
            "Gerado em " + datetime.now().strftime("%d/%m/%Y às %H:%M"),
            estilo_subtitulo,
        ),
        Spacer(1, 5 * mm),
    ])

    secao("Resumo executivo")
    if n_vendas:
        partes = [
            f"A Bookmarked regista {numero(n_vendas)} vendas concluídas "
            f"e {euro(receita)} de receita acumulada."
        ]
    else:
        partes = ["Ainda não há vendas concluídas para avaliar o desempenho comercial."]
    if not stock.empty:
        partes.append(
            f"Na análise de stock, {len(criticos)} títulos estão em estado crítico "
            f"e {len(atencao)} requerem atenção."
        )
    if variacao is not None and receita_prevista is not None:
        partes.append(
            "O modelo aponta para uma possível diminuição de receita."
            if variacao < 0 else
            "O modelo aponta para uma possível subida de receita."
            if variacao > 0 else
            "O modelo aponta para estabilidade da receita."
        )
    bloco(" ".join(partes), "Síntese")

    secao("Indicadores principais")
    indicadores = [
        ("Receita Total", euro(receita)),
        ("Número de Vendas", numero(n_vendas)),
        ("Unidades Vendidas", numero(unidades)),
        ("Ticket Médio", euro(ticket)),
    ]
    cards = Table([
        [Paragraph(nome, estilo_label) for nome, _ in indicadores],
        [Paragraph(valor, estilo_valor) for _, valor in indicadores],
    ], colWidths=[44 * mm] * 4, rowHeights=[10 * mm, 13 * mm])
    cards.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F4F7FB")),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#D8E0EA")),
        ("INNERGRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D8E0EA")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    elementos.append(cards)
    if n_vendas:
        bloco(
            f"O histórico contém {numero(n_vendas)} vendas concluídas, "
            f"{numero(unidades)} unidades vendidas e {euro(receita)} de receita. "
            f"O ticket médio é de {euro(ticket)} por venda. "
            "Estes valores são acumulados e, isoladamente, não demonstram "
            "se o negócio cresceu ou diminuiu face a outro período."
        )
    else:
        bloco("Não existem vendas concluídas para analisar nesta exportação.")

    secao("Previsão e qualidade do modelo")
    tabela_dados(["Indicador", "Valor"], [
        ["Mês previsto", periodo],
        ["Receita prevista", euro(receita_prevista)],
        ["Variação prevista", f"{decimal(variacao)}%" if variacao is not None else "—"],
        ["Tendência", seguro(metricas.get("tendencia", "—"))],
        ["Modelo", seguro(metricas.get("modelo", "—"))],
        ["MAE", euro(mae)],
        ["RMSE", euro(rmse)],
        ["R²", decimal(r2, 4)],
    ], [85 * mm, 91 * mm])

    if receita_prevista is not None:
        mensagem = f"O modelo estima {euro(receita_prevista)} de receita para {periodo}. "
        if proximo_periodo is not None and proximo_periodo == pd.Timestamp.now().to_period("M"):
            mensagem += (
                "Este mês já está em curso, pelo que a estimativa deve ser "
                "interpretada como uma projeção para o mês completo. "
            )
        if variacao is not None:
            direcao = "aumento" if variacao > 0 else "redução" if variacao < 0 else "estabilidade"
            mensagem += f"A variação estimada indica {direcao} ({decimal(variacao)}%). "
            if base_comparacao is not None:
                mensagem += (
                    f"Foi confirmada uma comparação com a receita real de "
                    f"{base_comparacao} ({euro(referencia)}). "
                )
            else:
                mensagem += (
                    "Não foi possível confirmar o período-base desta percentagem "
                    "a partir do histórico exportado; não deve ser automaticamente "
                    "interpretada como variação face ao mês anterior. "
                )
        if r2 is not None:
            qualidade = "limitada" if r2 < 0.5 else "moderada" if r2 < 0.8 else "elevada"
            mensagem += f"O R² de {decimal(r2, 4)} indica capacidade explicativa {qualidade} nos testes. "
        if mae is not None:
            mensagem += f"O erro absoluto médio (MAE) foi de {euro(mae)}. "
        mensagem += "Estas previsões apoiam a gestão, mas não garantem resultados futuros."
        bloco(mensagem)
    else:
        bloco("Não existem previsões disponíveis para este relatório.")

    elementos.append(PageBreak())
    secao("Top 5 livros mais vendidos")
    livro_lider = None
    if top_livros.empty:
        paragrafo("Não existem dados de vendas válidos para apresentar o ranking.")
    else:
        linhas = []
        for _, linha in top_livros.iterrows():
            linhas.append([
                Paragraph(seguro(linha["livro"]), estilo_celula),
                numero(inteiro(linha["quantidade"])),
            ])
        tabela_dados(["Livro", "Unidades Vendidas"], linhas, [136 * mm, 40 * mm])
        livro_lider = str(top_livros.iloc[0]["livro"])
        qtd_lider = inteiro(top_livros.iloc[0]["quantidade"])
        bloco(
            f'"{livro_lider}" lidera o ranking com {numero(qtd_lider)} unidades. '
            "O volume vendido ajuda a identificar a procura, mas não permite "
            "determinar o lucro sem considerar os custos de aquisição e operação."
        )

    secao("Stock e reposição")
    if stock.empty:
        paragrafo("Não existem dados de stock disponíveis.")
    else:
        paragrafo(
            f"Foram analisados {numero(len(stock))} livros. Existem "
            f"{numero(len(criticos))} títulos críticos e {numero(len(atencao))} "
            f"em atenção. A recomendação total é de {numero(total_repor)} "
            "unidades para reposição."
        )
        if not top_alertas.empty:
            linhas = []
            for _, linha in top_alertas.iterrows():
                linhas.append([
                    Paragraph(seguro(linha.get("titulo", "—")), estilo_celula),
                    seguro(linha.get("nivel", "—")),
                    numero(inteiro(linha.get("stock_disponivel"))),
                    decimal(valor_numero(linha.get("meses_cobertura"))),
                    numero(inteiro(linha.get("quantidade_recomendada"))),
                ])
            tabela_dados(
                ["Livro", "Nível", "Stock", "Cobertura (meses)", "Repor"],
                linhas, [73 * mm, 25 * mm, 18 * mm, 38 * mm, 22 * mm],
            )
        bloco(
            "A cobertura indica aproximadamente durante quantos meses o stock disponível "
            "poderá responder à procura média estimada. As quantidades sugeridas de "
            "reposição são indicativas e devem ser verificadas pelo gestor."
            if len(criticos) or len(atencao) else
            "Não há títulos classificados como Crítico ou Atenção nos dados disponíveis. "
            "Deve continuar a acompanhar-se a rotação de stock."
        )

    secao("Recomendações de gestão")
    recomendacoes = []
    if not criticos.empty:
        recomendacoes.append(
            "Dar prioridade aos títulos críticos: " + lista_titulos(criticos) + "."
        )
    if not atencao.empty:
        recomendacoes.append(
            "Acompanhar os títulos em atenção: " + lista_titulos(atencao) + "."
        )
    if livro_lider:
        recomendacoes.append(
            f'Verificar a disponibilidade de "{livro_lider}", por ser o livro mais vendido.'
        )
    if receita_prevista is not None:
        recomendacoes.append(
            "Confrontar a previsão de receita com os valores efetivamente registados "
            "e atualizar o modelo à medida que surgirem novos dados."
        )
    if not recomendacoes:
        recomendacoes.append(
            "Acompanhar regularmente as vendas, a disponibilidade de stock "
            "e as necessidades de reposição."
        )
    for i, recomendacao in enumerate(recomendacoes, 1):
        elementos.append(Paragraph(f"<b>{i}.</b> {seguro(recomendacao)}", estilo_normal))

    secao("Conclusão executiva")
    conclusao = (
        "O SABIN reúne indicadores operacionais e previsões que ajudam a orientar "
        "decisões comerciais. "
    )
    if len(criticos):
        conclusao += (
            "A prioridade imediata é avaliar a reposição dos títulos em estado crítico, "
            "considerando o stock disponível e a procura recente. "
        )
    elif not stock.empty:
        conclusao += "Não foram identificados títulos críticos nesta análise. "
    if variacao is not None and variacao < 0:
        conclusao += (
            "A possível descida prevista de receita justifica acompanhamento, "
            "mas deve ser confirmada com os resultados reais. "
        )
    conclusao += (
        "As recomendações são automáticas e não substituem a avaliação "
        "do responsável pela livraria."
    )
    bloco(conclusao, "Conclusão")
    elementos.append(Spacer(1, 2 * mm))
    paragrafo(
        "Relatório gerado automaticamente pelo SABIN. Os indicadores comerciais "
        "consideram apenas vendas concluídas. As previsões são estimativas e as "
        "recomendações de stock devem ser validadas antes de efetuar compras."
    )

    documento.build(elementos, onFirstPage=rodape_pdf, onLaterPages=rodape_pdf)
    return caminho


def encontrar_ficheiro_power_bi():
    caminho_env = os.getenv("POWER_BI_FILE")
    if caminho_env:
        caminho = Path(caminho_env)
        if not caminho.is_absolute():
            caminho = BASE_DIR / caminho
        if caminho.exists() and caminho.suffix.lower() == ".pbix":
            return caminho

    caminho_sabin = BASE_DIR / "powerbi" / "sabin.pbix"
    if caminho_sabin.exists():
        return caminho_sabin

    candidatos = []
    pasta_power_bi = BASE_DIR / "powerbi"
    if pasta_power_bi.exists():
        candidatos.extend(
            arquivo for arquivo in pasta_power_bi.iterdir()
            if arquivo.is_file() and arquivo.suffix.lower() == ".pbix"
        )
    candidatos.extend(
        arquivo for arquivo in BASE_DIR.iterdir()
        if arquivo.is_file() and arquivo.suffix.lower() == ".pbix"
    )
    if not candidatos:
        raise FileNotFoundError(
            "Não foi encontrado nenhum ficheiro Power BI (.pbix). "
            "Confirma se existe SABIN/powerbi/sabin.pbix."
        )
    return sorted(candidatos, key=lambda item: item.name.lower())[0]


def abrir_power_bi():
    caminho = encontrar_ficheiro_power_bi()
    if sys.platform.startswith("win"):
        os.startfile(str(caminho))
    elif sys.platform == "darwin":
        subprocess.Popen(["open", str(caminho)])
    else:
        subprocess.Popen(["xdg-open", str(caminho)])
    return caminho